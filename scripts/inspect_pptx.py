from pptx import Presentation

prs = Presentation('SIH2026-IDEA-Presentation-Format.pptx')
print(f"Slide Dimensions: {prs.slide_width.inches:.2f} x {prs.slide_height.inches:.2f} inches")
for idx, slide in enumerate(prs.slides):
    print(f"\n=================== SLIDE {idx+1} ===================")
    for s_idx, shape in enumerate(slide.shapes):
        txt = ""
        if shape.has_text_frame:
            txt = shape.text.strip().replace('\n', ' ')
            if len(txt) > 80:
                txt = txt[:80] + "..."
        print(f"[{s_idx}] Name: '{shape.name}', ID: {shape.shape_id}, L={shape.left/914400:.2f}\", T={shape.top/914400:.2f}\", W={shape.width/914400:.2f}\", H={shape.height/914400:.2f}\" | Text: '{txt}'")
